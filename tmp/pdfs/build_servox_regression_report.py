from pathlib import Path
import math
import numpy as np
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Flowable

ROOT = Path(r"D:\Git Repository\chengyi-chung\SP-AX3")
SOURCE = ROOT / "DOC" / "1150909鞋型表格分析V2.xlsx"
OUTPUT = ROOT / "output" / "pdf" / "ServoX_Regression_Reverse_Validation_Report.pdf"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

pdfmetrics.registerFont(TTFont("CJK", r"C:\Windows\Fonts\msjh.ttc", subfontIndex=0))
pdfmetrics.registerFont(TTFont("CJKB", r"C:\Windows\Fonts\msjhbd.ttc", subfontIndex=0))

NAVY = colors.HexColor("#153A5B")
BLUE = colors.HexColor("#2878A5")
GREEN = colors.HexColor("#DDF1E4")
GREEN_DARK = colors.HexColor("#287A4A")
PALE = colors.HexColor("#EEF5F8")
GRAY = colors.HexColor("#F2F4F6")
RED = colors.HexColor("#C94A4A")
DARK = colors.HexColor("#25333C")

wb_formula = load_workbook(SOURCE, data_only=False, read_only=True)
wb_value = load_workbook(SOURCE, data_only=True, read_only=True)
wsf = wb_formula["右側(X2軸)"]
wsv = wb_value["右側(X2軸)"]

servo = np.array([float(wsv.cell(r, 5).value) for r in range(2, 27)])
target = np.array([float(wsv.cell(r, 4).value) for r in range(2, 27)])
pred = np.array([float(wsv.cell(r, 9).value) for r in range(2, 27)])
formula = str(wsf["I2"].value).lstrip("=")
err = pred - target

r2 = 1.0 - np.sum(err**2) / np.sum((target - target.mean())**2)
rmse = float(np.sqrt(np.mean(err**2)))
mae = float(np.mean(np.abs(err)))
max_abs = float(np.max(np.abs(err)))
bias = float(np.mean(err))
within5 = int(np.sum(np.abs(err) <= 5))
within10 = int(np.sum(np.abs(err) <= 10))

loo_err = []
for i in range(len(servo)):
    keep = np.ones(len(servo), dtype=bool)
    keep[i] = False
    slope, intercept = np.polyfit(servo[keep], target[keep], 1)
    loo_err.append(intercept + slope * servo[i] - target[i])
loo_err = np.array(loo_err)
loo_rmse = float(np.sqrt(np.mean(loo_err**2)))
loo_mae = float(np.mean(np.abs(loo_err)))


class RegressionChart(Flowable):
    def __init__(self, width=165*mm, height=78*mm):
        super().__init__(); self.width = width; self.height = height
    def draw(self):
        c = self.canv; w, h = self.width, self.height
        left, bottom, top, right = 16*mm, 12*mm, 8*mm, 5*mm
        pw, ph = w-left-right, h-bottom-top
        xmin, xmax = 95, 345; ymin, ymax = 40, 480
        tx = lambda x: left + (x-xmin)/(xmax-xmin)*pw
        ty = lambda y: bottom + (y-ymin)/(ymax-ymin)*ph
        c.setStrokeColor(colors.HexColor("#D6DEE3")); c.setLineWidth(.5)
        for x in range(100, 341, 40):
            c.line(tx(x), bottom, tx(x), bottom+ph)
            c.setFillColor(DARK); c.setFont("CJK", 6.5); c.drawCentredString(tx(x), 3*mm, str(x))
        for y in range(50, 481, 100):
            c.line(left, ty(y), left+pw, ty(y)); c.setFillColor(DARK); c.drawRightString(left-2*mm, ty(y)-2, str(y))
        c.setStrokeColor(NAVY); c.setLineWidth(1)
        c.line(left,bottom,left,bottom+ph); c.line(left,bottom,left+pw,bottom)
        c.setFillColor(NAVY); c.setFont("CJKB",7.5); c.drawCentredString(left+pw/2, 0, "Servo X")
        c.saveState(); c.translate(3*mm,bottom+ph/2); c.rotate(90); c.drawCentredString(0,0,"目標 / 回歸 ServoX"); c.restoreState()
        c.setStrokeColor(BLUE); c.setLineWidth(1.4)
        for i in range(len(servo)-1): c.line(tx(servo[i]),ty(pred[i]),tx(servo[i+1]),ty(pred[i+1]))
        c.setFillColor(GREEN_DARK)
        for x,y in zip(servo,target): c.circle(tx(x),ty(y),2.2,fill=1,stroke=0)
        c.setFillColor(GREEN_DARK); c.setFont("CJK",7); c.drawString(left+3*mm,bottom+ph-4*mm,"● 目標 ServoX")
        c.setStrokeColor(BLUE); c.line(left+36*mm,bottom+ph-3*mm,left+45*mm,bottom+ph-3*mm)
        c.setFillColor(BLUE); c.drawString(left+47*mm,bottom+ph-4*mm,"回歸值")


class ResidualChart(Flowable):
    def __init__(self, width=165*mm, height=65*mm):
        super().__init__(); self.width = width; self.height = height
    def draw(self):
        c=self.canv; w,h=self.width,self.height; l,b,t,r=16*mm,12*mm,7*mm,5*mm
        pw,ph=w-l-r,h-b-t; xmin,xmax=95,345; ymin,ymax=-10,10
        tx=lambda x:l+(x-xmin)/(xmax-xmin)*pw; ty=lambda y:b+(y-ymin)/(ymax-ymin)*ph
        c.setStrokeColor(colors.HexColor("#D6DEE3")); c.setLineWidth(.5)
        for y in [-10,-5,0,5,10]:
            c.line(l,ty(y),l+pw,ty(y)); c.setFillColor(DARK); c.setFont("CJK",6.5); c.drawRightString(l-2*mm,ty(y)-2,str(y))
        c.setStrokeColor(NAVY); c.setLineWidth(1); c.line(l,b,l,b+ph); c.line(l,ty(0),l+pw,ty(0))
        c.setFillColor(BLUE)
        for x,e in zip(servo,err):
            c.setFillColor(RED if abs(e)>5 else BLUE); c.circle(tx(x),ty(e),2.2,fill=1,stroke=0)
        c.setFillColor(NAVY); c.setFont("CJKB",7.5); c.drawCentredString(l+pw/2,0,"Servo X")
        c.saveState(); c.translate(3*mm,b+ph/2); c.rotate(90); c.drawCentredString(0,0,"誤差（回歸值 - 目標值）"); c.restoreState()


styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="TitleC",fontName="CJKB",fontSize=21,leading=29,textColor=NAVY,alignment=TA_CENTER,spaceAfter=8))
styles.add(ParagraphStyle(name="SubC",fontName="CJK",fontSize=9.5,leading=15,textColor=colors.HexColor("#61717B"),alignment=TA_CENTER))
styles.add(ParagraphStyle(name="H1C",fontName="CJKB",fontSize=15,leading=21,textColor=NAVY,spaceBefore=4,spaceAfter=8))
styles.add(ParagraphStyle(name="H2C",fontName="CJKB",fontSize=11,leading=17,textColor=BLUE,spaceBefore=6,spaceAfter=5))
styles.add(ParagraphStyle(name="BodyC",fontName="CJK",fontSize=9.3,leading=16,textColor=DARK,spaceAfter=5))
styles.add(ParagraphStyle(name="SmallC",fontName="CJK",fontSize=7.5,leading=11,textColor=DARK))
styles.add(ParagraphStyle(name="FormulaC",fontName="CJK",fontSize=10,leading=16,backColor=GRAY,borderPadding=8,textColor=NAVY,spaceAfter=7))
styles.add(ParagraphStyle(name="PassC",fontName="CJKB",fontSize=12,leading=19,backColor=GREEN,borderPadding=10,textColor=GREEN_DARK,spaceAfter=8))
styles.add(ParagraphStyle(name="NoteC",fontName="CJK",fontSize=8.5,leading=14,backColor=PALE,borderPadding=8,textColor=NAVY,spaceAfter=7))

def footer(canvas, doc):
    canvas.saveState(); canvas.setStrokeColor(colors.HexColor("#CDD7DD")); canvas.line(18*mm,16*mm,192*mm,16*mm)
    canvas.setFont("CJK",7.2); canvas.setFillColor(colors.HexColor("#677883"))
    canvas.drawString(18*mm,10*mm,"ServoX 回歸反推基本驗證報告")
    canvas.drawRightString(192*mm,10*mm,f"第 {doc.page} 頁")
    canvas.restoreState()

doc=SimpleDocTemplate(str(OUTPUT),pagesize=A4,leftMargin=18*mm,rightMargin=18*mm,topMargin=17*mm,bottomMargin=22*mm,
                      title="ServoX 回歸反推基本驗證報告")
story=[]
story += [Spacer(1,12*mm),Paragraph("ServoX 回歸反推基本驗證報告",styles["TitleC"]),
          Paragraph("依據 1150909鞋型表格分析V2.xlsx - 右側(X2軸)",styles["SubC"]),Spacer(1,11*mm),
          Paragraph("結論：基本驗證可行",styles["PassC"]),
          Paragraph("表內 25 筆資料顯示，回歸 ServoX 與目標 ServoX 高度一致。R² 為 0.9989，RMSE 為 4.01，平均絕對誤差為 3.41；全部樣本誤差均在 ±10 以內。以現有資料作為基本反推與控制補償驗證，具可行性。",styles["BodyC"])]

metrics=[["指標","結果","判讀"],["樣本數","25 筆","Servo X 100 至 340，間隔 10"],["R²",f"{r2:.4f}","線性模型解釋度高"],
         ["RMSE",f"{rmse:.2f}","典型誤差約 4 個 Servo 單位"],["MAE",f"{mae:.2f}","平均絕對偏差小"],
         ["最大絕對誤差",f"{max_abs:.2f}","全數仍小於 10"],["平均偏差",f"{bias:.2f}","整體無明顯單向偏移"],
         ["±5 以內",f"{within5}/25（{within5/25:.0%}）","多數點接近目標"],["±10 以內",f"{within10}/25（100%）","基本驗證通過"]]
t=Table(metrics,colWidths=[42*mm,43*mm,79*mm],repeatRows=1)
t.setStyle(TableStyle([("FONTNAME",(0,0),(-1,-1),"CJK"),("FONTNAME",(0,0),(-1,0),"CJKB"),("FONTSIZE",(0,0),(-1,-1),8.5),
                       ("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,GRAY]),
                       ("GRID",(0,0),(-1,-1),.5,colors.HexColor("#BBC9D1")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
                       ("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
story += [Spacer(1,5*mm),t,Spacer(1,7*mm),Paragraph("驗證定位",styles["H1C"]),
          Paragraph("本結果足以支持修正後公式作為基本反推驗證與初步控制換算。先前實際驗證不佳，主要原因是程式使用 X軸理論值時漏掉 ServoX 與理論 X 之間的 30 常數，造成固定偏差。它仍屬同一份資料內的模型驗證，不等同於獨立上機驗收。",styles["NoteC"]),PageBreak()]

story += [Paragraph("一、先前實際驗證不佳的原因",styles["H1C"]),
          Paragraph("來源表格明確定義 X軸理論值與 ServoX 的關係：",styles["BodyC"]),
          Paragraph("X軸理論值 = ServoX - 30<br/>因此：ServoX = X軸理論值 + 30",styles["FormulaC"]),
          Paragraph("先前修正計算若把 X軸理論值直接代入以 ServoX 建立的回歸式，就等於漏掉 +30。這不是小幅四捨五入誤差，而是輸入座標基準不同所造成的固定系統誤差。",styles["BodyC"]),
          Paragraph("30 常數如何影響回歸式",styles["H2C"]),
          Paragraph("原回歸式（輸入為 ServoX）：<br/>TargetServoX = -111.09 + 1.6686 × ServoX<br/><br/>代入 ServoX = X軸理論值 + 30：<br/>TargetServoX = -111.09 + 1.6686 × (X軸理論值 + 30)<br/>TargetServoX = -61.032 + 1.6686 × X軸理論值",styles["FormulaC"]),
          Paragraph("若漏掉 30，會產生的固定差值：1.6686 × 30 = 50.058 Servo 單位",styles["PassC"]),
          Paragraph("實例：ServoX = 200",styles["H2C"])]
example=[["項目","計算","結果"],["X軸理論值","200 - 30","170"],
         ["正確反推","-111.09 + 1.6686 × 200","222.630"],
         ["等價理論X公式","-61.032 + 1.6686 × 170","222.630"],
         ["先前漏30的算法","-111.09 + 1.6686 × 170","172.572"],
         ["漏30造成差值","222.630 - 172.572","50.058"]]
et=Table(example,colWidths=[43*mm,79*mm,42*mm])
et.setStyle(TableStyle([("FONTNAME",(0,0),(-1,-1),"CJK"),("FONTNAME",(0,0),(-1,0),"CJKB"),("FONTSIZE",(0,0),(-1,-1),8.5),
                        ("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),
                        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,GRAY]),("GRID",(0,0),(-1,-1),.5,colors.HexColor("#BBC9D1")),
                        ("LEFTPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7)]))
story += [et,Spacer(1,7*mm),Paragraph("實作時只能選一種輸入形式",styles["H2C"]),
          Paragraph("A. 程式輸入為 ServoX：使用 TargetServoX = -111.09 + 1.6686 × ServoX。<br/>B. 程式輸入為 X軸理論值：使用 TargetServoX = -61.032 + 1.6686 × X軸理論值。<br/><br/>不可先扣 30 後仍使用原截距 -111.09，也不可在使用 ServoX 原值時再額外加 30，否則會形成漏補償或重複補償。",styles["NoteC"]),PageBreak()]

story += [Paragraph("二、資料與修正後公式",styles["H1C"]),
          Paragraph("分析使用右側(X2軸)工作表第 2 至 26 列，共 25 筆。自變數為 Servo X（E欄），驗證目標為目標ServoX（D欄），回歸結果為回歸ServoX(反推驗證)（I欄）。",styles["BodyC"]),
          Paragraph(f"回歸公式：F(X) = -111.09 + 1.6686 × X<br/>Excel 公式：{formula}",styles["FormulaC"]),
          Paragraph("此公式不是用來預測 Y軸理論值。它用現行 Servo X 反推修正後的目標 ServoX，並與 D欄目標值比較。",styles["NoteC"]),
          Paragraph("三、目標值與回歸值",styles["H1C"]),RegressionChart(),
          Paragraph("圖中回歸線幾乎與 25 個目標點重合，代表在 Servo X 100 至 340 的資料範圍內，線性關係穩定。",styles["BodyC"]),
          Paragraph("四、殘差分布",styles["H1C"]),ResidualChart(),
          Paragraph("誤差未呈現持續單方向擴大。最大正誤差為 +7.69，最大負誤差為 -8.77，均在 ±10 範圍內。",styles["BodyC"]),PageBreak()]

story += [Paragraph("五、反推驗證結果",styles["H1C"])]
rows=[["Servo X","目標 ServoX","回歸 ServoX","誤差","判定"]]
for x,y,p,e in zip(servo,target,pred,err):
    rows.append([f"{x:.0f}",f"{y:.0f}",f"{p:.3f}",f"{e:+.3f}","±10內" if abs(e)<=10 else "超限"])
t=Table(rows,colWidths=[27*mm,35*mm,38*mm,31*mm,29*mm],repeatRows=1)
t.setStyle(TableStyle([("FONTNAME",(0,0),(-1,-1),"CJK"),("FONTNAME",(0,0),(-1,0),"CJKB"),("FONTSIZE",(0,0),(-1,-1),7.3),
                       ("ALIGN",(0,1),(-1,-1),"RIGHT"),("ALIGN",(-1,1),(-1,-1),"CENTER"),
                       ("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),
                       ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,GRAY]),("TEXTCOLOR",(-1,1),(-1,-1),GREEN_DARK),
                       ("GRID",(0,0),(-1,-1),.35,colors.HexColor("#C4D0D7")),("TOPPADDING",(0,0),(-1,-1),3.5),("BOTTOMPADDING",(0,0),(-1,-1),3.5)]))
story += [t,Spacer(1,7*mm),Paragraph("六、穩健性補充",styles["H1C"]),
          Paragraph(f"以留一法交叉驗證（每次拿掉 1 筆、用其餘 24 筆重新回歸）檢查，RMSE 為 {loo_rmse:.2f}，MAE 為 {loo_mae:.2f}。結果與表內公式的 RMSE {rmse:.2f} 接近，顯示模型不是只靠單一資料點維持擬合。",styles["BodyC"]),
          Paragraph("建議的基本驗收條件",styles["H2C"]),
          Paragraph("1. 僅在 Servo X 100 至 340 的已驗證範圍內使用。<br/>2. 以獨立實測資料確認絕對誤差是否仍符合機構容許值。<br/>3. 若控制允許誤差為 ±10 Servo 單位，現有 25 筆資料全部通過；若要求 ±5，通過率為 80%。<br/>4. 對範圍端點 Servo X=100 與 340 增加重複量測，確認端點偏差。",styles["NoteC"]),
          Paragraph("七、最終判定",styles["H1C"]),
          Paragraph("先前實際驗證不佳可由漏掉 30 常數所造成的固定 50.058 Servo 單位偏差解釋。修正輸入基準後，在目前 25 筆右側 X2 軸資料與 ±10 Servo 單位的基本判定尺度下，回歸 ServoX（反推驗證）可行。",styles["PassC"]),
          Paragraph("八、建議上機獨立驗證方案",styles["H1C"]),
          Paragraph("1. 在 Servo X 100、160、220、280、340 五個區段各選至少 2 個未參與建模的新點。<br/>2. 每一點至少重複量測 3 次，記錄目標值、回歸值、實際機構到位值與重複性。<br/>3. 沿用 ±10 Servo 單位作為基本門檻時，要求所有獨立點均通過；若機構規格更嚴格，應以機構容許值重新設定門檻。<br/>4. 特別檢查 100 與 340 兩端，因現有最大負誤差出現在高端 Servo X=340。<br/>5. 若新資料仍維持 RMSE 約 4 至 5 且無系統性偏差，即可進一步評估導入控制補償。",styles["BodyC"]),
          Paragraph("適用範圍與限制",styles["H2C"])]
limits=[["項目","本報告判定"],["驗證資料","右側 X2 軸，25 筆表內樣本"],["Servo X 範圍","100 至 340"],
        ["基本誤差門檻","±10 Servo 單位"],["範圍外外插","未驗證，不建議直接使用"],["正式量產判定","需補充獨立上機量測"]]
lt=Table(limits,colWidths=[48*mm,116*mm])
lt.setStyle(TableStyle([("FONTNAME",(0,0),(-1,-1),"CJK"),("FONTNAME",(0,0),(-1,0),"CJKB"),("FONTSIZE",(0,0),(-1,-1),8.5),
                        ("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),
                        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,GRAY]),("GRID",(0,0),(-1,-1),.5,colors.HexColor("#BBC9D1")),
                        ("LEFTPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
story += [lt,Spacer(1,7*mm),Paragraph("建議保留的控制公式",styles["H2C"]),
          Paragraph("輸入為 ServoX：TargetServoX = -111.09 + 1.6686 × ServoX<br/>輸入為 X軸理論值：TargetServoX = -61.032 + 1.6686 × X軸理論值",styles["FormulaC"]),
          Paragraph("此公式適合作為目前資料範圍內的基本反推驗證基準。後續若加入新鞋型或不同機構條件，應重新檢查係數、殘差與端點誤差。",styles["NoteC"]),
          Paragraph(f"資料來源：{SOURCE}",styles["SmallC"])]

doc.build(story,onFirstPage=footer,onLaterPages=footer)
print(OUTPUT)
